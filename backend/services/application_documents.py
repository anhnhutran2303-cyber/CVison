"""Session-only application documents. No provider calls or invented CV facts."""
from io import BytesIO
import re
from typing import Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator

from backend.exceptions import InvalidCVError
from backend.models import CVDocument
from backend.normalization import extract_known_skills, normalized_text, valid_evidence
from backend.parsers.cv_parser import extract_sections
from backend.parsers.jd_parser import JDParser
from backend.services.bullet_review_service import section_bullets


SECTIONS = {
    "en": {"summary": "Summary", "experience": "Experience", "projects": "Projects",
           "education": "Education", "skills": "Skills", "certifications": "Certifications"},
    "vi": {"summary": "Tóm tắt", "experience": "Kinh nghiệm", "projects": "Dự án",
           "education": "Học vấn", "skills": "Kỹ năng", "certifications": "Chứng chỉ"},
}
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
BULLET = re.compile(r"^[-•*]\s+")


class DocumentInput(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)

    @field_validator("*", mode="before")
    @classmethod
    def safe_text(cls, value):
        if isinstance(value, str) and INVALID_XML.search(value):
            raise ValueError("Remove unsupported control characters.")
        return value


class ResumeDraft(DocumentInput):
    full_name: str = Field(min_length=1, max_length=120)
    headline: str = Field(default="", max_length=200)
    email: str = Field(default="", max_length=200)
    phone: str = Field(default="", max_length=80)
    location: str = Field(default="", max_length=200)
    links: str = Field(default="", max_length=1000)
    summary: str = Field(default="", max_length=4000)
    experience: str = Field(default="", max_length=15000)
    projects: str = Field(default="", max_length=10000)
    education: str = Field(default="", max_length=6000)
    skills: str = Field(default="", max_length=4000)
    certifications: str = Field(default="", max_length=4000)
    language: Literal["en", "vi"] = "en"
    template: Literal["classic", "modern"] = "modern"

    @field_validator("email")
    @classmethod
    def email_shape(cls, value):
        if value and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value):
            raise ValueError("Enter a valid email address.")
        return value

    @model_validator(mode="after")
    def has_content(self):
        if not any(getattr(self, field) for field in SECTIONS[self.language]):
            raise ValueError("Add at least one CV section.")
        return self


def resume_text(draft: ResumeDraft) -> str:
    header = [draft.full_name, draft.headline,
              " | ".join(value for value in (draft.email, draft.phone, draft.location) if value), draft.links]
    parts = ["\n".join(value for value in header if value)]
    for field, title in SECTIONS[draft.language].items():
        if content := getattr(draft, field):
            parts.append(f"{title}\n{content}")
    return "\n\n".join(parts)


def _document(template="modern"):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    font = "Calibri" if template == "modern" else "Cambria"
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "List Bullet"):
        style = doc.styles[name]
        style.font.name = font
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        fonts = style.element.rPr.rFonts
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attr}"), None)
        fonts.set(qn("w:eastAsia"), font)
        fonts.set(qn("w:cs"), font)
    # Some python-docx distributions include a decorative border in Title.
    # Remove inherited borders too so both exports stay plain across runtimes.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Title"].paragraph_format.space_after = Pt(4)
    doc.styles["Heading 1"].font.size = Pt(12)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    return doc


def _paragraphs(doc, text):
    for line in text.splitlines():
        line = line.strip()
        if line:
            marked = bool(BULLET.match(line))
            doc.add_paragraph(BULLET.sub("", line) if marked else line,
                              style="List Bullet" if marked else "Normal")


def _bytes(doc):
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def resume_docx(draft: ResumeDraft) -> bytes:
    doc = _document(draft.template)
    title = doc.add_paragraph(draft.full_name, style="Title")
    if draft.template == "classic":
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if draft.headline:
        doc.add_paragraph(draft.headline, style="Subtitle")
    contact = " | ".join(value for value in (draft.email, draft.phone, draft.location) if value)
    if contact:
        doc.add_paragraph(contact)
    _paragraphs(doc, draft.links)
    for field, heading in SECTIONS[draft.language].items():
        if content := getattr(draft, field):
            doc.add_paragraph(heading, style="Heading 1")
            _paragraphs(doc, content)
    return _bytes(doc)


def evidence_options(cv_text: str, jd_text: str = "") -> list[str]:
    """Rank genuine CV excerpts by known skills in assessable JD sections."""
    if not cv_text.strip() or len(cv_text) > 100000:
        return []
    candidates = [quote for _, quote in section_bullets(CVDocument(raw_text=cv_text))]
    for section, content in extract_sections(cv_text).items():
        if section in {"experience", "projects", "education", "skills", "summary"}:
            candidates.extend(BULLET.sub("", line.strip()) for line in content.splitlines()
                              if 15 <= len(line.strip()) <= 1000)
    unique = {}
    for quote in candidates:
        if len(quote) <= 1500 and valid_evidence([quote], cv_text):
            unique.setdefault(normalized_text(quote), quote)
    jd_skills = set(extract_known_skills("\n".join(line for line, _ in JDParser().assessable_lines(jd_text))))
    return sorted(unique.values(), key=lambda quote: -len(jd_skills & set(extract_known_skills(quote))))[:30]


class LetterRequest(DocumentInput):
    full_name: str = Field(min_length=1, max_length=120)
    role: str = Field(min_length=1, max_length=200)
    company: str = Field(min_length=1, max_length=200)
    recipient: str = Field(default="", max_length=200)
    motivation: str = Field(default="", max_length=2000)
    cv_text: str = Field(min_length=1, max_length=100000)
    jd_text: str = Field(min_length=1, max_length=100000)
    evidence: list[str] = Field(min_length=1, max_length=3)
    language: Literal["en", "vi"] = "en"

    @model_validator(mode="after")
    def grounded_evidence(self):
        offered = set(evidence_options(self.cv_text, self.jd_text))
        if len(set(self.evidence)) != len(self.evidence) or any(q not in offered for q in self.evidence):
            raise ValueError("Choose one to three different excerpts from the current CV.")
        return self


def cover_letter(request: LetterRequest) -> str:
    quotes = "\n".join(f"- {quote}" for quote in request.evidence)
    if request.language == "vi":
        parts = [f"Kính gửi {request.recipient or 'Ban Tuyển dụng'},",
                 f"Tôi viết thư này để ứng tuyển vị trí {request.role} tại {request.company}.",
                 request.motivation,
                 "Tôi xin chia sẻ những nội dung sau từ CV của mình:\n" + quotes,
                 "Tôi mong có cơ hội trao đổi thêm về kinh nghiệm của mình và yêu cầu của vị trí. "
                 "Cảm ơn Anh/Chị đã dành thời gian xem xét hồ sơ.",
                 f"Trân trọng,\n{request.full_name}"]
    else:
        parts = [f"Dear {request.recipient or 'Hiring Team'},",
                 f"I am applying for the {request.role} position at {request.company}.",
                 request.motivation,
                 "I would like to highlight the following from my CV:\n" + quotes,
                 "I would welcome the opportunity to discuss my experience and the requirements of the role. "
                 "Thank you for considering my application.",
                 f"Sincerely,\n{request.full_name}"]
    return "\n\n".join(part for part in parts if part)


def letter_docx(text: str) -> bytes:
    if not text.strip() or len(text) > 20000 or INVALID_XML.search(text):
        raise InvalidCVError("Enter a letter of 1–20,000 characters without unsupported control characters.")
    doc = _document()
    for block in text.split("\n\n"):
        _paragraphs(doc, block)
        if doc.paragraphs:
            doc.paragraphs[-1].paragraph_format.space_after = Pt(12)
    return _bytes(doc)
