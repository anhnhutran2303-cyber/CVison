from io import BytesIO

def extract_text_from_upload(uploaded_file):
    """Extract text from TXT, PDF, or DOCX. Returns (text, error)."""
    if uploaded_file is None:
        return "", None
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    try:
        if name.endswith(".txt"):
            try:
                return data.decode("utf-8"), None
            except UnicodeDecodeError:
                return data.decode("latin-1"), None
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
            if not text.strip():
                return "", "Không đọc được nội dung chữ từ PDF này. Hãy thử DOCX/TXT hoặc dán nội dung CV."
            return text, None
        if name.endswith(".docx"):
            from docx import Document
            doc = Document(BytesIO(data))
            chunks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        chunks.append(row_text)
            text = "\n".join(chunks)
            if not text.strip():
                return "", "Không đọc được nội dung chữ từ DOCX này."
            return text, None
        return "", "Chỉ hỗ trợ PDF, DOCX và TXT."
    except Exception as exc:
        return "", f"Không thể đọc file: {exc}"
